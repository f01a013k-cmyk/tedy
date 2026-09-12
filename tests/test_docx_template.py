"""既存の Word 様式への流し込み.

支援機関は自前のフォーマットを持っている。こちらの体裁で出し直すと使ってもらえない
ため、相手の .docx をそのまま開いて中身だけ埋める。書式が壊れないことが要件。
"""

from __future__ import annotations

import pytest
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from jizokuka.render.docx_template import fill, heading_key, normalize
from jizokuka.samples import load_shien_fill, shien_form_path


def _body(doc: Document) -> list[tuple[str, object]]:
    out = []
    for ch in doc.element.body.iterchildren():
        if ch.tag.endswith("}p"):
            out.append(("p", Paragraph(ch, doc)))
        elif ch.tag.endswith("}tbl"):
            out.append(("t", Table(ch, doc)))
    return out


@pytest.fixture
def filled(tmp_path):
    blocks, tables = load_shien_fill()
    path, missing = fill(shien_form_path(), tmp_path / "out.docx", blocks, tables)
    assert missing == []
    return Document(path), blocks, tables


# -- 見出しの解釈 -----------------------------------------------------------


def test_全角の章番号を解釈する():
    assert heading_key("１－１－１　設立経緯（社名の由来）") == (
        "1-1-1",
        "設立経緯(社名の由来)",
    )


def test_長音記号をハイフンに変えない():
    """正規化しすぎると「顧客ニーズ」が「顧客ニ-ズ」になり照合が壊れる."""
    assert heading_key("２－２　顧客ニーズ") == ("2-2", "顧客ニーズ")
    assert "サービス" in heading_key("３－２　自社の提供する商品・サービスの強み")[1]


def test_見出しでない段落はNoneを返す():
    assert heading_key("例：代表　氏名　××××") is None
    assert heading_key("①") is None
    assert heading_key("") is None


@pytest.mark.parametrize(
    "line",
    [
        "2026年度（初年度）　売上高3,700万円（うちEC売上200万円）、営業利益180万円",
        "1,996,800円 × 54% = 約107万円（初年度の粗利増加額）",
        "4月　冷凍機の機種選定と発注、電源増設工事の手配",
        "2年目以降は、初年度購入者のリピート率を30%と見込む。",
        "① 餡を内製しており、糖度を1度刻みに調整できる",
        "・店頭で発送を希望しながら断ってきた分　月20件",
        "7,000,000円 × 54% = 378万円",
    ],
)
def test_数字で始まる本文を見出しと誤認しない(line):
    """章番号の後には必ず空白が入る。ここを緩めると流し込み先を取り違える.

    丸数字は NFKC で「①」→「1」になるため、正規化前に弾く必要がある。
    """
    assert heading_key(line) is None


def test_全角スラッシュを正規化する():
    """表の列見出し「販売個数／日」を半角表記のキーで引けること."""
    assert normalize("販売個数／日") == "販売個数/日"


# -- 流し込みの結果 ---------------------------------------------------------


def test_全ての項目が文書に入る(filled):
    doc, blocks, _ = filled
    text = "".join(p.text for p in doc.paragraphs)
    for key, content in blocks.items():
        head = content.strip().split("\n")[0][:20]
        assert head in text, f"{key} の本文が入っていない"


def test_章番号が重複する項目も取り違えない(filled):
    """様式には 2-2 が2つ、3-2 が2つある（採番の重複）."""
    doc, _, _ = filled
    text = "".join(p.text for p in doc.paragraphs)
    assert "松本市の人口は" in text          # 2-2 市場の動向
    assert "顧客からは具体的な要望" in text   # 2-2 顧客ニーズ
    assert "原材料表示が1品あたり7品目以内" in text  # 3-2 商品の強み
    assert "冷凍・結露に対応しておらず" in text      # 3-2 商品の弱み


def test_記入例のプレースホルダは残らない(filled):
    doc, _, _ = filled
    text = "".join(p.text for p in doc.paragraphs)
    for placeholder in ("××××", "●●●●", "●●機", "〇名"):
        assert placeholder not in text, f"記入例 {placeholder} が残っている"


def test_様式の指示文は残す(filled):
    """「〜を挙げてください」は様式の一部なので消さない."""
    doc, _, _ = filled
    text = "".join(p.text for p in doc.paragraphs)
    assert "戦略と戦術をどう考えているかを挙げてください。" in text
    assert "部門別に集計している場合には" in text


def test_空欄が残らない(filled):
    doc, _, _ = filled
    assert [p for p in doc.paragraphs if not p.text.strip()] == []


def test_本文は表より前に置かれる(filled):
    """「下表のとおりである」と書いた段落が表の後ろに回らないこと."""
    doc, _, _ = filled
    body = _body(doc)
    intro = next(
        i for i, (k, o) in enumerate(body)
        if k == "p" and "直近3期の実績と今期予想は下表のとおり" in o.text
    )
    table_at = next(i for i, (k, _) in enumerate(body) if k == "t" and i > intro)
    assert intro < table_at


def test_記入欄が無い見出しでも本文が入る(filled):
    """7-1 は見出しの直後がいきなり表で、空欄が1つも無い."""
    doc, _, _ = filled
    text = "".join(p.text for p in doc.paragraphs)
    assert "客単価3,200円 × 新規購入者52人/月 × 12か月" in text


def test_表が全て埋まる(filled):
    doc, _, tables = filled
    assert len(doc.tables) == 5
    # 決算推移表
    kessan = doc.tables[1]
    assert kessan.rows[1].cells[1].text == "4,200"
    assert kessan.rows[4].cells[4].text == "180"
    # 全角スラッシュの列も埋まっていること
    assert doc.tables[0].rows[1].cells[3].text == "172個"


def test_様式の空行に追加項目を書き込める(filled):
    """競合比較表の「他は自由に項目を挙げてください」行を使う."""
    doc, _, _ = filled
    labels = [r.cells[0].text for r in doc.tables[2].rows]
    assert "主力商品" in labels and "餡" in labels and "規模" in labels
    assert "他は自由に項目を挙げてください" not in labels


def test_セル内の改行はWordの改行になる(filled):
    doc, _, _ = filled
    cell = doc.tables[3].rows[1].cells[1]
    assert cell._tc.xml.count("<w:br/>") >= 2


def test_様式に無い見出しを指定すると落ちる(tmp_path):
    with pytest.raises(ValueError, match="見出しが見つかりません"):
        fill(shien_form_path(), tmp_path / "x.docx", {"9-9-9": "本文"}, {})


def test_原本は書き換えない(filled):
    """テンプレートは読み取り専用で扱う."""
    doc = Document(shien_form_path())
    blanks = [p for p in doc.paragraphs if not p.text.strip()]
    assert blanks, "原本の空欄が消えている"
