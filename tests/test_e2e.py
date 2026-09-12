"""パイプライン全体の結合テスト（LLM はスタブ）."""

from __future__ import annotations

import pytest

from jizokuka.generator import build
from jizokuka.render import docx_form, markdown, report, xlsx_expense


@pytest.fixture
def built(sheet, fake_llm, koubo):
    return build(sheet, fake_llm, koubo=koubo, max_revisions=2, target_score=80,
                 report=lambda m: None)


def test_全セクションが生成される(built, koubo):
    plan = built.plan
    keys = {s.key for s in plan.all_sections()}
    assert keys == {spec["key"] for spec in koubo.all_section_specs()}
    assert all(s.body.strip() for s in plan.all_sections())


def test_事業名は30字以内に収まる(built):
    name = built.plan.section("project_name")
    assert name.char_count <= 30
    assert not name.over_limit


def test_採点が低ければリライトが走りスコアが上がる(built):
    assert len(built.score_history) >= 2
    assert built.score_history[-1] > built.score_history[0]
    assert built.plan.score.total >= 80


def test_資金調達が経費合計と一致する(built):
    total = sum(e.resolved_amount() for e in built.plan.expenses)
    assert sum(built.plan.funding.values()) == total
    assert built.plan.funding["補助金"] == 333_333  # 500,000 × 2/3 切捨


def test_インボイス特例が上限に反映される(built, koubo):
    assert built.plan.specials == ["インボイス特例"]
    assert koubo.limit_yen(built.plan.frame, built.plan.specials) == 1_000_000


def test_準拠チェックが実行されている(built):
    assert built.plan.compliance is not None
    rules = {i.rule for i in built.plan.compliance.issues}
    assert "koubo.verified" in rules  # 未検証ナレッジの警告は必ず出る


def test_ギャップ分析の質問が保持される(built):
    assert built.gap.questions
    assert built.gap.questions[0].choices


def test_全ての成果物ファイルが生成できる(built, koubo, tmp_path):
    plan = built.plan
    paths = [
        markdown.write(tmp_path / "plan.md", markdown.plan_markdown(plan, koubo)),
        markdown.write(tmp_path / "q.md", markdown.questions_markdown(built.gap, "t")),
        markdown.write(tmp_path / "score.md", report.score_report(plan, built.score_history)),
        docx_form.build_yoshiki2(plan, koubo, tmp_path / "y2.docx"),
        docx_form.build_yoshiki3(plan, koubo, tmp_path / "y3.docx"),
        xlsx_expense.build(plan, koubo, tmp_path / "expense.xlsx"),
    ]
    for p in paths:
        assert p.exists() and p.stat().st_size > 0, p


def test_生成したdocxが読み戻せる(built, koubo, tmp_path):
    from docx import Document

    docx_form.build_yoshiki2(built.plan, koubo, tmp_path / "y2.docx")
    text = "\n".join(p.text for p in Document(tmp_path / "y2.docx").paragraphs)
    assert "経営計画書兼補助事業計画書①" in text
    assert "企業概要" in text
    assert "補助事業の効果" in text


def test_生成したxlsxの補助金が数式になっている(built, koubo, tmp_path):
    from openpyxl import load_workbook

    xlsx_expense.build(built.plan, koubo, tmp_path / "e.xlsx")
    wb = load_workbook(tmp_path / "e.xlsx")
    assert wb["経費明細表"]["E2"].value == "=C2*D2"
    grant = wb["資金調達方法"]["B7"].value
    assert grant.startswith("=MIN(ROUNDDOWN(")


def test_スコアレポートに改善指示が載る(built):
    md = report.score_report(built.plan, built.score_history)
    assert "審査観点スコアレポート" in md
    assert "改善指示" in md or "指摘なし" in md


def test_質問シートは選択肢形式で出る(built):
    md = markdown.questions_markdown(built.gap, "t")
    assert "- [ ]" in md
    assert "その他" in md


def test_同梱のヒアリングシート雛形がそのままパースできる():
    """インストール後も `jizokuka template` が動くこと（パッケージデータ）."""
    import yaml

    from jizokuka.config import TEMPLATE_DIR
    from jizokuka.models import HearingSheet

    path = TEMPLATE_DIR / "hearing_sheet.yaml"
    assert path.exists(), f"雛形がパッケージに含まれていません: {path}"
    sheet = HearingSheet.model_validate(yaml.safe_load(path.read_text(encoding="utf-8")))
    assert sheet.frame == "通常枠"
    assert sheet.total_expense() == 0  # 空欄のままでも壊れない
