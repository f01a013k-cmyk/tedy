"""Excel 出力 — 経費明細表・資金調達方法.

支援員が金額を調整すると補助金額が自動再計算されるよう、
固定値ではなく**数式**を埋め込む。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..knowledge import Koubo
from ..models import Plan

_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
_WARN_FILL = PatternFill("solid", fgColor="FFE5E5")
_YEN = '#,##0"円"'


def _style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, size=10)
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build(plan: Plan, koubo: Koubo, out_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "経費明細表"

    headers = ["経費区分", "内容・必要理由", "単価(税抜)", "数量", "補助対象経費", "積算根拠", "対応する取組"]
    widths = [16, 32, 12, 7, 14, 34, 26]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[get_column_letter(i)].width = w
    _style_header(ws, 1, len(headers))

    valid_categories = set(koubo.category_names())
    first_data_row = 2
    row = first_data_row
    for e in plan.expenses:
        ws.cell(row=row, column=1, value=e.category or "")
        ws.cell(row=row, column=2, value=e.item or "")
        if e.unit_price:
            ws.cell(row=row, column=3, value=e.unit_price).number_format = _YEN
        if e.quantity:
            ws.cell(row=row, column=4, value=e.quantity)
        # 単価×数量が入っていれば数式にし、支援員の編集に追随させる
        if e.unit_price and e.quantity:
            ws.cell(row=row, column=5, value=f"=C{row}*D{row}")
        else:
            ws.cell(row=row, column=5, value=e.resolved_amount())
        ws.cell(row=row, column=5).number_format = _YEN
        ws.cell(row=row, column=6, value=e.basis or "")
        ws.cell(row=row, column=7, value=e.linked_activity or "")

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=c)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        # 対象外の疑い・区分不正はセルを着色して目視で気づけるようにする
        problems = koubo.ineligible_hits(f"{e.item or ''} {e.basis or ''}")
        if problems or (e.category not in valid_categories):
            for c in (1, 2):
                ws.cell(row=row, column=c).fill = _WARN_FILL
            note = (
                "；".join(f"{kw}: {reason}" for kw, reason in problems)
                or f"経費区分『{e.category}』は補助対象経費区分にありません"
            )
            ws.cell(row=row, column=2).comment = _comment(note)
        row += 1

    last_data_row = row - 1
    ws.cell(row=row, column=1, value="合計").font = Font(bold=True)
    total_cell = ws.cell(
        row=row, column=5, value=f"=SUM(E{first_data_row}:E{last_data_row})"
    )
    total_cell.font = Font(bold=True)
    total_cell.number_format = _YEN
    for c in range(1, len(headers) + 1):
        ws.cell(row=row, column=c).border = _BORDER
    total_ref = f"'{ws.title}'!E{row}"

    _funding_sheet(wb, plan, koubo, total_ref)
    _reference_sheet(wb, koubo)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _comment(text: str):
    from openpyxl.comments import Comment

    return Comment(text, "jizokuka")


def _funding_sheet(wb: Workbook, plan: Plan, koubo: Koubo, total_ref: str) -> None:
    ws = wb.create_sheet("資金調達方法")
    rate = koubo.rate(plan.frame)
    limit = koubo.limit_yen(plan.frame, plan.specials)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 44

    rows = [
        ("申請枠", plan.frame, ""),
        ("上乗せ特例", "・".join(plan.specials) or "なし", ""),
        ("補助率", koubo.rate_label(plan.frame), f"={rate.numerator}/{rate.denominator}"),
        ("補助上限額", limit, "枠＋特例の合計"),
        ("補助対象経費 合計", f"={total_ref}", "経費明細表の合計と連動"),
    ]
    for i, (k, v, note) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=v)
        if isinstance(v, int):
            c.number_format = _YEN
        ws.cell(row=i, column=3, value=note).font = Font(size=9, color="777777")

    # 補助金申請額 = MIN(経費×率, 上限) を数式で持たせる
    ws.cell(row=7, column=1, value="補助金交付申請額").font = Font(bold=True)
    grant = ws.cell(
        row=7,
        column=2,
        value=f"=MIN(ROUNDDOWN(B5*{rate.numerator}/{rate.denominator},0),B4)",
    )
    grant.number_format = _YEN
    grant.font = Font(bold=True)
    ws.cell(row=7, column=3, value="経費または上限を変えると自動で再計算されます").font = Font(
        size=9, color="777777"
    )

    ws.cell(row=9, column=1, value="【資金調達方法】").font = Font(bold=True)
    ws.cell(row=10, column=1, value="区分").font = Font(bold=True)
    ws.cell(row=10, column=2, value="金額").font = Font(bold=True)
    _style_header(ws, 10, 2)

    ws.cell(row=11, column=1, value="補助金")
    ws.cell(row=11, column=2, value="=B7").number_format = _YEN
    ws.cell(row=12, column=1, value="自己資金")
    ws.cell(row=12, column=2, value="=B5-B7").number_format = _YEN
    ws.cell(row=13, column=1, value="合計").font = Font(bold=True)
    tot = ws.cell(row=13, column=2, value="=SUM(B11:B12)")
    tot.number_format = _YEN
    tot.font = Font(bold=True)
    ws.cell(row=14, column=3, value="※合計は補助対象経費合計(B5)と一致する必要があります").font = Font(
        size=9, color="CC0000"
    )
    for r in range(10, 14):
        for c in (1, 2):
            ws.cell(row=r, column=c).border = _BORDER


def _reference_sheet(wb: Workbook, koubo: Koubo) -> None:
    """経費区分と注意事項の早見表. 支援員が手で直すときに参照する."""
    ws = wb.create_sheet("補助対象経費リファレンス")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 52

    for i, h in enumerate(["No", "経費区分", "例", "注意事項"], start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws, 1, 4)

    for r, c in enumerate(koubo.expense_categories, start=2):
        ws.cell(row=r, column=1, value=c["code"])
        ws.cell(row=r, column=2, value=c["name"])
        ws.cell(row=r, column=3, value="、".join(c["examples"]))
        ws.cell(row=r, column=4, value=c.get("caution") or "")
        for col in range(1, 5):
            ws.cell(row=r, column=col).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=r, column=col).border = _BORDER

    start = len(koubo.expense_categories) + 4
    ws.cell(row=start, column=1, value="【補助対象外】").font = Font(bold=True, color="CC0000")
    r = start + 1
    for rule in koubo.raw.get("ineligible_keywords", []):
        ws.cell(row=r, column=2, value="、".join(rule["keyword"]))
        ws.cell(row=r, column=3, value=rule["reason"])
        ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True)
        r += 1
